import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
D = ROOT / "docs" / "diagrams"
S = ROOT / "docs" / "evidence" / "screens"
nfr = json.loads((ROOT / "docs" / "evidence" / "nfr_results.json").read_text())
ui = json.loads((ROOT / "docs" / "evidence" / "ui_walkthrough.json").read_text())
NAVY = RGBColor(0x1E, 0x3A, 0x8A)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin = sec.bottom_margin = Inches(0.8)
base = doc.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(10.5)
base.paragraph_format.space_after = Pt(4)
for name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = NAVY
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.keep_with_next = True

figure_no = [0]
table_no = [0]


def para(text, bold=False, italic=False, size=None, align=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.add_run(item)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def caption(text, kind):
    counter = figure_no if kind == "Figure" else table_no
    counter[0] += 1
    p = para(f"{kind} {counter[0]}. {text}", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    return p


def table(headers, rows, widths, size=8.5, title=None):
    if title:
        p = caption(title, "Table")
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(size)
        r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "1E3A8A")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(value))
            r.font.size = Pt(size)
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(path, width, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    caption(text, "Figure")


def side_by_side(items):
    t = doc.add_table(rows=2, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (path, width, text) in enumerate(items):
        c = t.rows[0].cells[i]
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].add_run().add_picture(str(path), width=Inches(width))
        figure_no[0] += 1
        cap = t.rows[1].cells[i].paragraphs[0]
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figure {figure_no[0]}. {text}")
        r.italic, r.font.size = True, Pt(9)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------- Title page ----------------
for _ in range(4):
    para("")
para("NMIMS Mumbai — Mukesh Patel School of Technology Management & Engineering", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
para("B.Tech Computer Engineering — Software Engineering", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
para("Software Engineering Mini Project Report", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
t = para("Student Recommendation Platform for New City Students", bold=True, size=24, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
t.runs[0].font.color.rgb = NAVY
table(["Project author", ""], [
    ["Name", "Krishang Vikram Singh"], ["Roll number", "B106"], ["Batch", "C1"],
    ["Programme", "B.Tech Computer Engineering"], ["Subject", "Software Engineering"],
    ["Other group members", "[to be filled in if applicable]"], ["Faculty / guide", "[to be filled in]"],
    ["Date of submission", "[to be filled in]"]], [2.0, 3.6])
para("")
para("Source code, tests, diagrams and evidence are in the project folder student-recommendation-platform (see README.md).",
     italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
page_break()

# ---------------- 1 Introduction ----------------
doc.add_heading("1. Introduction, problem statement and objectives", 1)
para("Students who move to a new city for college must quickly find a place to live, somewhere to eat, a stationery shop, "
     "a place to play sport and things to do. Advice comes from scattered word of mouth, and online reviews are anonymous "
     "and hard to trust. The Student Recommendation Platform is a web application where verified students of a college "
     "rate and review such places, and an administrator moderates what is published. It was developed across ten Software "
     "Engineering experiments (requirements, use cases, activity, class, sequence, collaboration and data-flow design, testing "
     "and corrective coding) and this report presents the resulting working system.")
doc.add_heading("Problem statement", 2)
para("New students lack a trustworthy, location-aware source of information about accommodation, food, stationery, sport "
     "and events near their campus, and existing review sites do not confirm that a reviewer is a student of the college or "
     "check what is published.")
doc.add_heading("Objectives", 2)
bullets(["Let students register with a college email and log in securely, with lockout after repeated failures.",
         "Let anyone search nearby places by category, filter by cost, rating and distance, and browse upcoming events.",
         "Let only verified students submit 1–5 star reviews, and publish them only after administrator approval, so that average ratings reflect approved reviews only.",
         "Give administrators tools to moderate reviews and reports and to manage students, colleges, cities and listings, with an audit trail.",
         "Verify the system with automated tests, a live smoke test, and measurements against the quality targets of Experiment 2."])
doc.add_heading("Scope", 2)
para("In scope: the five categories (PGs/hostels, restaurants, stationery shops, sports facilities, events), college-domain "
     "verification, reviews, reports, notifications, personalised recommendations and the administrator functions above. "
     "Out of scope (not built): HTTPS hosting, email delivery, maps or third-party business data, college-ID verification, "
     "slot booking and other Experiment 5 operations that have no user story. All listings in the demonstration are sample data.")

# ---------------- 2 Requirements ----------------
doc.add_heading("2. Requirements summary", 1)
para("Requirements come from Experiment 1 (eleven Agile user stories, US1–US11) and Experiment 2 (quality targets). "
     "Experiment 10 added four business rules found through testing. The full inventory, traceability matrix and a list of "
     "17 contradictions between the experiments are in the docs folder.")
table(["ID", "Functional requirement (Experiment 1)", "Status"], [
    ["US1–US2", "Register with college email; duplicate and non-college emails rejected; account marked verified", "Done; domain check only"],
    ["US3", "Login, error message, lock for 15 minutes after 5 consecutive failures", "Done"],
    ["US4, US9", "Search within 5 km sorted by distance; filters for cost, rating, distance; empty-result guidance", "Done"],
    ["US5", "Events in date order with type and date filters", "Done"],
    ["US6–US7", "Verified students review and report; reports make the review “Under Review”; no duplicate reports", "Done; publication after approval"],
    ["US8", "Personalised recommendations from saved preferences and location", "Done"],
    ["US10", "Notifications with read state and per-category opt-out", "Done (in-app)"],
    ["US11", "Admin moderation, suspension and reinstatement, logged", "Done; no edit-flag action"]],
    [0.8, 4.6, 1.3], title="Functional requirements and their status")
table(["Quality attribute", "Target (Experiment 2)"], [
    ["Performance", "Non-search ≤ 1 s; search ≤ 3 s"], ["Availability", "99.5 % per month"],
    ["Security", "Authenticated, encrypted sessions; failed attempts limited and logged"],
    ["Retention / reliability", "Reviews kept ≥ 2 years; 99.9 % of submissions recorded correctly"],
    ["Usability", "Register and search in ≤ 5 clicks"], ["Recoverability", "Service restored within 8 h"],
    ["Scalability", "Search ≤ 3 s with 100 000 users and 500 concurrent sessions"],
    ["Confidentiality, integrity", "No private details exposed; invalid ratings, empty fields and spam rejected"],
    ["Portability, accessibility", "Chrome/Edge/Safari and phone sizes; WCAG 2.1 AA"]],
    [1.7, 5.0], title="Non-functional targets")

# ---------------- 3 Architecture and design ----------------
doc.add_heading("3. System architecture and design", 1)
para("The system is a small layered monolith (Figure 1). A browser single-page interface calls a JSON API implemented with "
     "the Python standard library HTTP server. All business rules live in one service class that is independent of HTTP. "
     "Working data is held in memory and every change is written through to a SQLite database whose schema enforces keys, "
     "unique emails and the 1–5 rating range. Python was chosen because the earlier experiments already described Python logic "
     "and pytest tests; it keeps the project dependency-free and easy to demonstrate.")
figure(D / "architecture.png", 5.6, "Architecture of the implemented system")
table(["Module", "Responsibility"], [
    ["web/ (HTML, CSS, JavaScript)", "Search, events, recommendations, notifications, account and administrator screens"],
    ["server.py", "Routing, bearer-token sessions with 30-minute idle timeout, role checks, safe error messages, security headers"],
    ["service.py", "Registration, login lockout, search, recommendations, reviews, moderation, notifications, administration, audit log"],
    ["models.py, errors.py, security.py", "Domain objects, error types, scrypt password hashing"],
    ["store.py", "SQLite schema (users, colleges, cities, recommendations, reviews, reports, notifications, audit log) and write-through saves"]],
    [2.1, 4.6], title="Modules")
doc.add_heading("3.1 Use cases and class design", 2)
para("The use-case diagram (Figure 2) shows the Student and Administrator actors. The original Experiment 3 image was not "
     "available for this build, so the diagram was redrawn from Experiments 1, 4 and 5 and is labelled as reconstructed. "
     "The class diagram (Figure 3) is the one drawn in Experiment 5: an abstract User with Student and Administrator, an "
     "abstract Recommendation with five subclasses, College, City and Review.")
figure(D / "use_case.png", 4.9, "Use-case diagram (reconstructed from Experiments 1, 4 and 5)")
figure(D / "class_diagram_exp5.png", 6.3, "Class diagram from Experiment 5")
para("Differences between the class diagram and the code are listed in the consistency analysis. In short: one users table with a "
     "role column replaces the two user subclasses; the category-specific attributes (rent range, cuisine, timings, sport types, "
     "organizer, venue) are stored in a JSON details column; Notification, Report and Audit entries were added because "
     "US7, US10 and Experiment 2 need them; operations such as bookSlot and updateMenu have no user story and were not built.")
doc.add_heading("3.2 Workflow and interaction design", 2)
para("Figure 4 (activity diagram, Experiment 4) and Figure 5 (sequence diagram, Experiment 6) show the main path: register, "
     "verify, log in, search, view details, submit a review, administrator validation and notification.")
figure(D / "activity_diagram_exp4.png", 4.3, "Activity diagram with Student and Administrator swimlanes (Experiment 4)")
figure(D / "sequence_2.png", 4.5, "Sequence diagram: search, view, submit and validate a review (Experiment 6)")
doc.add_heading("3.3 Data flow", 2)
para("The Level 0 and Level 1 data-flow diagrams (Experiment 8) show four processes — manage accounts, search and recommend, "
     "manage reviews, administer platform data — and four data stores (users, colleges and cities, recommendations, reviews). "
     "These correspond to the registration/login, search, review and administration methods of the service class and the "
     "matching tables in the database.")
figure(D / "dfd_level1.png", 6.0, "Data flow diagram, Level 1 (Experiment 8)")

# ---------------- 4 Implementation ----------------
doc.add_heading("4. Implementation", 1)
para("The code is about 1 400 lines of Python and a 260-line JavaScript interface, with no third-party runtime packages. Key rules:")
bullets(["Registration: the email domain must equal the college domain or be a subdomain of it, so “notnmims.edu.in” is rejected; "
         "emails are lower-cased before comparison; passwords need 8 or more characters and are stored as salted scrypt hashes.",
         "Login: five consecutive failures lock the account for 15 minutes; the lock expires by itself; suspended accounts are refused "
         "and their open sessions are ended; every failure is written to the audit log.",
         "Search: haversine distance, default radius 5 km, filters for cost and rating, results sorted by distance and limited to 50; "
         "events are hidden once their date has passed.",
         "Reviews: integer rating 1–5 (booleans and fractions rejected), mandatory comment up to 1 000 characters, URL/repeat/profanity "
         "screening, one review per student per listing; reviews start Pending and count towards the average only when Published; "
         "a reported review leaves the average until the administrator resolves it.",
         "Administration: every administrator method checks the role on the server; catalogue edits, verification, suspension and "
         "moderation are logged with actor, time and target; listings are archived rather than deleted so reviews are retained.",
         "Confidentiality: public responses contain display names and review text only; the administrator list shows masked emails."])
para("A limitation worth stating plainly: verification is by email domain. The platform sends no confirmation email, so it cannot prove "
     "that the student owns the mailbox; an administrator can verify students manually.")

# ---------------- 5 Testing ----------------
doc.add_heading("5. Testing methodology and results", 1)
para("Testing combined pytest unit and API tests, a live smoke test, a browser walkthrough and load measurements. The first "
     "prototype (Experiment 9) failed 4 of 26 tests; the fixes of Experiment 10 made all 26 pass, and no test was weakened afterwards. "
     "The suite was then extended to 96 tests.")
table(["Check", "Command / tool", "Result"], [
    ["Automated tests", "python -m pytest -v", "96 passed, 0 failed; 92 % line coverage of src/srp"],
    ["Core cases TC-01 to TC-04", "tests/test_case_study.py", "All pass: non-college email rejected; lock after 5 failures; unverified student blocked; approved reviews give 4.5"],
    ["Historical defects D-1 to D-4", "tests/test_extended.py, test_features.py", "All pass: look-alike domain, case-different duplicate, suspended login, 4.5 and True ratings rejected"],
    ["Live smoke test", "scripts/smoke_check.py", "18 checks passed against the running server (201/422/409/403/423/401 statuses as expected)"],
    ["Restart", "stop and start on the same database", "Accounts, lock state, review and rating preserved"],
    ["Browser walkthrough", "Playwright, Chromium", "Register, search, review, approve and recommend flow completed; 0 axe-core violations on 7 screens"]],
    [1.6, 2.1, 3.0], title="Test summary")
doc.add_heading("Quality targets: measured results", 2)
cs, seq, lr, rs, rr = nfr["concurrent_search"], nfr["sequential"], nfr["login_ramp"], nfr["review_submission"], nfr["restart_recovery"]
table(["Attribute", "Measured (one 4-core machine, loopback)", "Verdict"], [
    ["Performance", f"Search median {seq['search']['median_s']*1000:.0f} ms, login median {seq['login']['median_s']*1000:.0f} ms (100 001 users, 20 000 reviews)", "Met for single requests"],
    ["Burst of logins", f"20 parallel logins: median {lr['median_s']:.2f} s, max {lr['max_s']:.2f} s (hashing runs one at a time)", "1 s target missed in this burst"],
    ["Scalability", f"500 sessions, 500 simultaneous searches: all succeeded, max {cs['max_s']:.2f} s", "Met for search"],
    ["Reliability", f"{rs['http_201']}/{rs['attempted']} simultaneous reviews stored, {rs['rows_in_database']} database rows", "Met in this test"],
    ["Recoverability", f"Restart healthy in {rr['seconds_to_healthy']} s with all data", "Process restart only"],
    ["Usability", f"{ui['clicks_register_and_search']} clicks to register and search", "Met"],
    ["Accessibility / portability", "axe-core: 0 violations; Chromium desktop and 390 px phone only", "Partly met"],
    ["Availability, encrypted transport", "Not measured; plain HTTP only", "Not verified / not met"]],
    [1.5, 3.9, 1.3], title="Experiment 2 targets")
para("These figures describe this environment only. Availability and long-term retention cannot be established by tests run in a "
     "single session; HTTPS was not implemented; portability and accessibility were checked only with Chromium and an automated scanner.",
     italic=True, size=9.5)

# ---------------- 6 Screens ----------------
doc.add_heading("6. Screenshots and demonstration results", 1)
para("The screens below were captured from the running application by the walkthrough script (sample data).")
side_by_side([(S / "03_search_results.png", 3.2, "Search results with rating, cost, distance and review count"),
              (S / "07_admin_queue.png", 2.6, "Administrator dashboard with the pending review")])
para("")
side_by_side([(S / "09_rating_after_approval.png", 3.1, "After approval the rating and review count update"),
              (S / "10_events.png", 3.1, "Events in chronological order")])

# ---------------- 7 Deployment ----------------
doc.add_heading("7. Deployment", 1)
para("The application runs as one process: set PYTHONPATH=src and SRP_ADMIN_PASSWORD, then run python -m srp. It listens on "
     "http://127.0.0.1:8000/, creates data/platform.db and loads sample data on first run. This local deployment was started, "
     "exercised by the smoke test and restarted to confirm that data persists. No public hosting exists; the README explains how to "
     "add a TLS proxy before any public use.")

# ---------------- 8 Limitations ----------------
doc.add_heading("8. Limitations and future enhancements", 1)
bullets(["No HTTPS; email ownership and college ID are not verified (confirmation links would fix the first).",
         "Publication needs administrator approval, so a review does not appear within one minute as US6 words it.",
         "“Edit-flag” moderation, slot booking, menu and inventory updates, push or email notifications and real maps are not built.",
         "Single process, in-memory sessions, no backup or archival job; login bursts exceeded 1 s because password hashing is serialised.",
         "Availability, long-term retention, Safari/Edge/Firefox and manual screen-reader testing remain unverified.",
         "Future work: confirmation emails, TLS deployment, a hosted database, per-IP rate limiting, cached search index, and a user study."])

# ---------------- 9 Conclusion ----------------
doc.add_heading("9. Conclusion", 1)
para("The project turned the requirements and designs of Experiments 1–8 into a working, persistent and tested web application. "
     "All eleven user stories are implemented except where noted, the four defects found in testing are fixed and guarded by regression "
     "tests, and 96 automated tests pass. Measured performance meets the search and scalability targets on one machine, while "
     "encrypted transport, availability, broader browser coverage and manual accessibility checks are still open. The traceability "
     "matrix and consistency analysis make clear which claims are backed by evidence and which are not.")

# ---------------- References ----------------
doc.add_heading("References", 1)
for ref in [
    "K. V. Singh, Software Engineering Lab Record, Experiments 1–10 (B106, Batch C1), NMIMS MPSTME, 2026 (requirements, use case, activity, class, sequence, collaboration, data-flow, testing and coding experiments).",
    "M. Cohn, User Stories Applied: For Agile Software Development, Addison-Wesley, 2004.",
    "M. Fowler, UML Distilled: A Brief Guide to the Standard Object Modeling Language, 3rd ed., Addison-Wesley, 2003.",
    "C. Gane and T. Sarson, Structured Systems Analysis: Tools and Techniques, Prentice-Hall, 1979.",
    "I. Sommerville, Software Engineering, 10th ed., Pearson, 2016.",
    "R. W. Sinnott, “Virtues of the Haversine,” Sky and Telescope, vol. 68, no. 2, p. 159, 1984.",
    "W3C, Web Content Accessibility Guidelines (WCAG) 2.1, https://www.w3.org/TR/WCAG21/",
    "Python Software Foundation, Python 3 documentation: sqlite3, http.server, hashlib.scrypt, https://docs.python.org/3/",
    "pytest documentation, https://docs.pytest.org/ ; Deque Systems, axe-core, https://github.com/dequelabs/axe-core ; Playwright, https://playwright.dev/"]:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(ref)
    r.font.size = Pt(9.5)

# footer page numbers
for sct in doc.sections:
    fp = sct.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Krishang Vikram Singh · B106 · Student Recommendation Platform · page ")
    run.font.size = Pt(8)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = fp.add_run()
        r.font.size = Pt(8)
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)

out = ROOT / "report" / "Software_Engineering_Mini_Project_Report.docx"
doc.core_properties.author = "Krishang Vikram Singh"
doc.core_properties.title = "Student Recommendation Platform for New City Students"
doc.save(out)
print(out)
